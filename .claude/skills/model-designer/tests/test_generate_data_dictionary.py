# SPDX-License-Identifier: Apache-2.0
# Licensed to the Ed-Fi Alliance under one or more agreements.
# The Ed-Fi Alliance licenses this file to you under the Apache License, Version 2.0.
# See the LICENSE and NOTICES files in the project root for more information.

import sys
import unittest
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
from generate_data_dictionary import (  # noqa: E402
    CORE_FIT_VALUES,
    CRITIQUE_MODE,
    REFERENCE_MERGED_DOC_HEIGHT,
    DESIGN_MODE,
    build_workbook,
    load_input,
    _lines_needed,
)

SAMPLE_ENTITIES = [
    {
        "name": "StaffResponsibilityAssociation",
        "documentation": "Staff responsibility association.",
        "properties": [
            {
                "name": "EducationOrganization",
                "datatype": "Reference",
                "other": "MetaEd DSL Type: domain entity",
                "documentation": "The ed-org this responsibility is held at.",
                "identity": "Yes",
                "cardinality": "",
                "core_fit": "Maps cleanly",
            },
            {
                "name": "ResponsibilityDescriptor",
                "datatype": "Reference",
                "other": "MetaEd DSL Type: descriptor",
                "documentation": "The kind of responsibility held.",
                "identity": "",
                "cardinality": "required",
                "core_fit": "New to core",
            },
        ],
    },
    {
        "name": "StaffEdOrgEmploymentAssociation",
        "documentation": "Existing employment association.",
        "properties": [
            {
                "name": "EmploymentPeriod",
                "datatype": "Reference",
                "other": "MetaEd DSL Type: common",
                "documentation": "Now a collection to carry status history.",
                "identity": "",
                "cardinality": "required collection",
                "core_fit": "Partially maps",
            },
            {
                "name": "HireDate",
                "datatype": "Date",
                "other": "MetaEd DSL Type: date",
                "documentation": "Unchanged for context.",
                "identity": "",
                "cardinality": "required",
                "core_fit": "Established core concept, new placement",
            },
        ],
    },
]


def _property(**overrides):
    prop = {
        "name": "P",
        "datatype": "String",
        "other": "",
        "documentation": "",
        "identity": "",
        "cardinality": "",
        "core_fit": "Maps cleanly",
    }
    prop.update(overrides)
    return prop


class TestBuildWorkbook(unittest.TestCase):
    def setUp(self):
        self.output_path = Path(__file__).resolve().parent / "_tmp_test_output.xlsx"
        if self.output_path.exists():
            self.output_path.unlink()

    def tearDown(self):
        if self.output_path.exists():
            self.output_path.unlink()

    def test_one_sheet_per_entity_with_expected_title(self):
        build_workbook(SAMPLE_ENTITIES, str(self.output_path))
        wb = openpyxl.load_workbook(self.output_path)
        self.assertEqual(
            wb.sheetnames,
            ["StaffResponsibilityAssociation", "StaffEdOrgEmploymentAssociation"],
        )

    def test_header_row_is_reference_six_columns_plus_core_fit(self):
        build_workbook(SAMPLE_ENTITIES, str(self.output_path))
        wb = openpyxl.load_workbook(self.output_path)
        ws = wb["StaffResponsibilityAssociation"]
        header_row = [cell.value for cell in ws[4]]
        self.assertEqual(
            header_row,
            [
                "Property Name",
                "UML Datatype",
                "Other",
                "Documentation",
                "Identity",
                "Cardinality",
                "Core Fit",
            ],
        )

    def test_entity_name_and_documentation_in_first_two_rows(self):
        build_workbook(SAMPLE_ENTITIES, str(self.output_path))
        wb = openpyxl.load_workbook(self.output_path)
        ws = wb["StaffEdOrgEmploymentAssociation"]
        self.assertEqual(ws["A1"].value, "StaffEdOrgEmploymentAssociation")
        self.assertEqual(ws["A2"].value, "Existing employment association.")

    def test_property_rows_written_in_order(self):
        build_workbook(SAMPLE_ENTITIES, str(self.output_path))
        wb = openpyxl.load_workbook(self.output_path)
        ws = wb["StaffEdOrgEmploymentAssociation"]
        row5 = [cell.value for cell in ws[5]]
        row6 = [cell.value for cell in ws[6]]
        self.assertEqual(row5[0], "EmploymentPeriod")
        self.assertEqual(row5[5], "required collection")
        self.assertEqual(row6[0], "HireDate")
        self.assertEqual(row6[5], "required")

    def test_empty_entities_list_raises_value_error(self):
        with self.assertRaises(ValueError):
            build_workbook([], str(self.output_path))

    def test_missing_property_key_raises_value_error_naming_the_key(self):
        prop = _property()
        del prop["other"]
        bad_entities = [
            {
                "name": "StaffResponsibilityAssociation",
                "documentation": "doc",
                "properties": [prop],
            }
        ]
        with self.assertRaises(ValueError) as ctx:
            build_workbook(bad_entities, str(self.output_path))
        message = str(ctx.exception)
        self.assertIn("StaffResponsibilityAssociation", message)
        self.assertIn("other", message)

    def test_missing_entity_key_raises_value_error_naming_the_key(self):
        bad_entities = [
            {
                "name": "StaffResponsibilityAssociation",
                # "documentation" key deliberately missing
                "properties": [],
            }
        ]
        with self.assertRaises(ValueError) as ctx:
            build_workbook(bad_entities, str(self.output_path))
        message = str(ctx.exception)
        self.assertIn("StaffResponsibilityAssociation", message)
        self.assertIn("documentation", message)

    def test_long_entity_name_truncates_sheet_title_but_keeps_full_name_in_cell(self):
        long_name = "StudentEducationOrganizationResponsibilityAssociation"
        entities = [{"name": long_name, "documentation": "doc", "properties": [_property()]}]
        build_workbook(entities, str(self.output_path))
        wb = openpyxl.load_workbook(self.output_path)
        self.assertEqual(wb.sheetnames, [long_name[:31]])
        ws = wb[long_name[:31]]
        self.assertEqual(ws["A1"].value, long_name)


class TestRowHeightAutofit(unittest.TestCase):
    """Data rows must carry NO explicit height, so Excel autofits them.

    A fixed chars-per-width estimate cannot reproduce Excel's per-glyph font
    metrics. Calibrated against all 83 data rows of the reference workbook,
    every constant either clipped text or over-padded; the clipped case was a
    real defect, where a two-line Other cell carrying a shared string plus its
    SQL Recommended DataType got 29.0pt where the reference gives 43.5pt,
    hiding the SQL line entirely.
    """

    def setUp(self):
        self.output_path = Path(__file__).resolve().parent / "_tmp_autofit.xlsx"
        if self.output_path.exists():
            self.output_path.unlink()

    def tearDown(self):
        if self.output_path.exists():
            self.output_path.unlink()

    def _shared_string_entity(self):
        return [
            {
                "name": "StudentApplication",
                "documentation": "An application submitted on behalf of a student.",
                "properties": [
                    _property(
                        name="ApplicationIdentifier",
                        other=(
                            "MetaEd DSL Type: shared string\n"
                            "SQL Recommended DataType: VARCHAR(20)"
                        ),
                        documentation=(
                            "A unique number or alphanumeric code assigned to an application."
                        ),
                        identity="Yes",
                    )
                ],
            }
        ]

    def test_data_rows_have_no_explicit_height(self):
        build_workbook(self._shared_string_entity(), str(self.output_path))
        wb = openpyxl.load_workbook(self.output_path)
        ws = wb["StudentApplication"]
        self.assertIsNone(ws.row_dimensions[5].height)

    def test_data_rows_are_not_marked_custom_height(self):
        build_workbook(self._shared_string_entity(), str(self.output_path))
        wb = openpyxl.load_workbook(self.output_path)
        ws = wb["StudentApplication"]
        self.assertFalse(ws.row_dimensions[5].customHeight)

    def test_other_column_wraps_so_autofit_can_expand_the_row(self):
        build_workbook(self._shared_string_entity(), str(self.output_path))
        wb = openpyxl.load_workbook(self.output_path)
        ws = wb["StudentApplication"]
        self.assertTrue(ws.cell(row=5, column=3).alignment.wrap_text)
        self.assertIn("VARCHAR(20)", ws.cell(row=5, column=3).value)

    def test_merged_documentation_row_keeps_explicit_height(self):
        # Excel does not autofit merged cells, so row 2 must stay explicit.
        build_workbook(self._shared_string_entity(), str(self.output_path))
        wb = openpyxl.load_workbook(self.output_path)
        ws = wb["StudentApplication"]
        self.assertIsNotNone(ws.row_dimensions[2].height)

    def test_lines_needed_counts_explicit_newlines_as_separate_lines(self):
        text = "MetaEd DSL Type: shared string\nSQL Recommended DataType: VARCHAR(20)"
        # 30 chars, then 37 chars, at width 35 -> 1 + 2 = 3 lines.
        self.assertEqual(_lines_needed(text, 35.0), 3)

    def test_lines_needed_handles_empty_and_single_line(self):
        self.assertEqual(_lines_needed("", 35.0), 1)
        self.assertEqual(_lines_needed("MetaEd DSL Type: descriptor", 35.0), 1)


class TestCoreFitColumn(unittest.TestCase):
    def setUp(self):
        self.output_path = Path(__file__).resolve().parent / "_tmp_corefit.xlsx"
        if self.output_path.exists():
            self.output_path.unlink()

    def tearDown(self):
        if self.output_path.exists():
            self.output_path.unlink()

    def test_core_fit_values_are_the_four_extension_fit_categories(self):
        self.assertEqual(
            CORE_FIT_VALUES,
            (
                "Maps cleanly",
                "Partially maps",
                "Established core concept, new placement",
                "New to core",
            ),
        )

    def test_core_fit_written_to_seventh_column(self):
        build_workbook(SAMPLE_ENTITIES, str(self.output_path))
        wb = openpyxl.load_workbook(self.output_path)
        ws = wb["StaffResponsibilityAssociation"]
        self.assertEqual(ws.cell(row=5, column=7).value, "Maps cleanly")
        self.assertEqual(ws.cell(row=6, column=7).value, "New to core")

    def test_missing_core_fit_key_raises_value_error_naming_the_key(self):
        prop = _property()
        del prop["core_fit"]
        bad = [{"name": "E", "documentation": "doc", "properties": [prop]}]
        with self.assertRaises(ValueError) as ctx:
            build_workbook(bad, str(self.output_path))
        self.assertIn("core_fit", str(ctx.exception))

    def test_invalid_core_fit_value_raises_value_error_naming_value_and_property(self):
        bad = [
            {
                "name": "E",
                "documentation": "doc",
                "properties": [_property(name="Widget", core_fit="Brand New")],
            }
        ]
        with self.assertRaises(ValueError) as ctx:
            build_workbook(bad, str(self.output_path))
        message = str(ctx.exception)
        self.assertIn("Brand New", message)
        self.assertIn("Widget", message)

    def test_missing_is_not_an_allowed_core_fit_value(self):
        # "Missing" describes an element absent from the design, which by
        # definition has no row in this workbook.
        self.assertNotIn("Missing", CORE_FIT_VALUES)
        bad = [
            {
                "name": "E",
                "documentation": "doc",
                "properties": [_property(core_fit="Missing")],
            }
        ]
        with self.assertRaises(ValueError):
            build_workbook(bad, str(self.output_path))

    def test_blank_core_fit_is_allowed_for_unresolved_rows(self):
        entities = [
            {
                "name": "E",
                "documentation": "doc",
                "properties": [_property(name="SchoolChoice", core_fit="")],
            }
        ]
        build_workbook(entities, str(self.output_path))
        wb = openpyxl.load_workbook(self.output_path)
        # An empty-string cell round-trips as None, the same way the blank
        # Identity and Cardinality cells do -- the point is that it saved at
        # all rather than raising.
        self.assertIsNone(wb["E"].cell(row=5, column=7).value)

    def test_verify_marker_is_allowed_as_core_fit(self):
        entities = [
            {
                "name": "E",
                "documentation": "doc",
                "properties": [_property(name="SchoolChoice", core_fit="[Verify]")],
            }
        ]
        build_workbook(entities, str(self.output_path))
        wb = openpyxl.load_workbook(self.output_path)
        self.assertEqual(wb["E"].cell(row=5, column=7).value, "[Verify]")


class TestCommonsSheet(unittest.TestCase):
    """Commons get one shared tab, laid out like the entity sheets.

    A common's shape belongs in rows a reader can scan, not prose crammed
    into the Other column of the row that references it.
    """

    def setUp(self):
        self.output_path = Path(__file__).resolve().parent / "_tmp_commons.xlsx"
        if self.output_path.exists():
            self.output_path.unlink()

    def tearDown(self):
        if self.output_path.exists():
            self.output_path.unlink()

    def _entity_plus_two_commons(self):
        return [
            {
                "name": "StudentApplication",
                "documentation": "An application submitted for a student.",
                "properties": [
                    _property(name="ApplicantIdentificationCode", datatype="Reference")
                ],
            },
            {
                "kind": "common",
                "name": "ApplicantIdentificationCode",
                "documentation": "An identification code for the applicant.",
                "properties": [
                    _property(
                        name="StudentIdentificationSystem",
                        datatype="Reference",
                        identity="Yes",
                    ),
                    _property(name="IdentificationCode", cardinality="required"),
                ],
            },
            {
                "kind": "common",
                "name": "SchoolChoice",
                "documentation": "Shape not yet provided by TEA.",
                "properties": [],
            },
        ]

    def test_commons_land_on_one_sheet_named_commons(self):
        build_workbook(self._entity_plus_two_commons(), str(self.output_path))
        wb = openpyxl.load_workbook(self.output_path)
        self.assertEqual(wb.sheetnames, ["StudentApplication", "Commons"])

    def test_no_commons_sheet_when_there_are_no_commons(self):
        build_workbook(SAMPLE_ENTITIES, str(self.output_path))
        wb = openpyxl.load_workbook(self.output_path)
        self.assertNotIn("Commons", wb.sheetnames)

    def test_kind_defaults_to_entity(self):
        # SAMPLE_ENTITIES carries no "kind" key at all.
        build_workbook(SAMPLE_ENTITIES, str(self.output_path))
        wb = openpyxl.load_workbook(self.output_path)
        self.assertEqual(len(wb.sheetnames), len(SAMPLE_ENTITIES))

    def test_unknown_kind_raises_value_error_naming_the_kind(self):
        bad = [{"kind": "widget", "name": "E", "documentation": "d", "properties": []}]
        with self.assertRaises(ValueError) as ctx:
            build_workbook(bad, str(self.output_path))
        self.assertIn("widget", str(ctx.exception))

    def test_each_common_gets_its_own_title_band_and_header(self):
        build_workbook(self._entity_plus_two_commons(), str(self.output_path))
        ws = openpyxl.load_workbook(self.output_path)["Commons"]
        col_a = [c.value for c in ws["A"]]
        self.assertIn("ApplicantIdentificationCode", col_a)
        self.assertIn("SchoolChoice", col_a)
        # The entity sheets' header, repeated once per block.
        self.assertEqual(col_a.count("Property Name"), 2)

    def test_common_property_rows_follow_their_header(self):
        build_workbook(self._entity_plus_two_commons(), str(self.output_path))
        ws = openpyxl.load_workbook(self.output_path)["Commons"]
        rows = [[c.value for c in row] for row in ws.iter_rows()]
        header_idx = next(i for i, r in enumerate(rows) if r[0] == "Property Name")
        self.assertEqual(rows[header_idx + 1][0], "StudentIdentificationSystem")
        self.assertEqual(rows[header_idx + 2][0], "IdentificationCode")

    def test_commons_sheet_uses_the_same_header_as_entity_sheets(self):
        build_workbook(self._entity_plus_two_commons(), str(self.output_path))
        wb = openpyxl.load_workbook(self.output_path)
        entity_header = [c.value for c in wb["StudentApplication"][4]]
        ws = wb["Commons"]
        rows = [[c.value for c in row] for row in ws.iter_rows()]
        common_header = next(r for r in rows if r[0] == "Property Name")
        self.assertEqual(common_header, entity_header)

    def test_common_with_no_properties_still_renders_its_block(self):
        build_workbook(self._entity_plus_two_commons(), str(self.output_path))
        ws = openpyxl.load_workbook(self.output_path)["Commons"]
        col_a = [c.value for c in ws["A"]]
        idx = col_a.index("SchoolChoice")
        self.assertEqual(col_a[idx + 1], "Shape not yet provided by TEA.")

    def test_commons_data_rows_have_no_explicit_height(self):
        build_workbook(self._entity_plus_two_commons(), str(self.output_path))
        ws = openpyxl.load_workbook(self.output_path)["Commons"]
        rows = [[c.value for c in row] for row in ws.iter_rows()]
        header_idx = next(i for i, r in enumerate(rows) if r[0] == "Property Name")
        data_row = header_idx + 2  # 1-indexed row just after the header
        self.assertIsNone(ws.row_dimensions[data_row].height)

    def test_commons_blocks_are_separated_by_a_blank_row(self):
        build_workbook(self._entity_plus_two_commons(), str(self.output_path))
        ws = openpyxl.load_workbook(self.output_path)["Commons"]
        col_a = [c.value for c in ws["A"]]
        second_title = col_a.index("SchoolChoice")
        self.assertIsNone(col_a[second_title - 1])



class TestCritiqueModeInvariant(unittest.TestCase):
    """A critique workbook renders what was SUBMITTED, not what is recommended.

    An element the submission does not contain must not have a row. A blank
    Core Fit is the tell: every submitted element is classifiable, so the only
    reason a critique row has no category is that it was never submitted.
    This is the defect that put an `ApplicationStatus` row in an as-submitted
    workbook whose own narrative doc classed it Missing.
    """

    def setUp(self):
        self.output_path = Path(__file__).resolve().parent / "_tmp_mode_output.xlsx"
        if self.output_path.exists():
            self.output_path.unlink()

    def tearDown(self):
        if self.output_path.exists():
            self.output_path.unlink()

    def _entities(self, core_fit):
        return [
            {
                "name": "StudentApplication",
                "documentation": "As submitted.",
                "properties": [_property(name="ApplicationStatus", core_fit=core_fit)],
            }
        ]

    def test_critique_mode_rejects_blank_core_fit(self):
        with self.assertRaises(ValueError) as ctx:
            build_workbook(self._entities(""), str(self.output_path), mode=CRITIQUE_MODE)
        self.assertIn("ApplicationStatus", str(ctx.exception))

    def test_critique_mode_error_explains_the_row_should_not_exist(self):
        with self.assertRaises(ValueError) as ctx:
            build_workbook(self._entities(""), str(self.output_path), mode=CRITIQUE_MODE)
        message = str(ctx.exception)
        self.assertIn("submission", message)
        self.assertIn("critique", message)

    def test_critique_mode_writes_nothing_when_it_rejects(self):
        with self.assertRaises(ValueError):
            build_workbook(self._entities(""), str(self.output_path), mode=CRITIQUE_MODE)
        self.assertFalse(self.output_path.exists())

    def test_design_mode_still_allows_blank_core_fit(self):
        # The blank escape stays legal in design mode -- an honest blank beats
        # a guessed label. Only the as-submitted render forbids it.
        build_workbook(self._entities(""), str(self.output_path), mode=DESIGN_MODE)
        self.assertTrue(self.output_path.exists())

    def test_design_mode_is_the_default_for_direct_calls(self):
        build_workbook(self._entities(""), str(self.output_path))
        self.assertTrue(self.output_path.exists())

    def test_critique_mode_allows_verify_placeholder(self):
        # [Verify] means "submitted, shape unresolved" -- legitimate in a
        # critique. Only the blank is forbidden.
        build_workbook(self._entities("[Verify]"), str(self.output_path), mode=CRITIQUE_MODE)
        self.assertTrue(self.output_path.exists())

    def test_critique_mode_allows_every_real_category(self):
        for value in CORE_FIT_VALUES:
            with self.subTest(core_fit=value):
                if self.output_path.exists():
                    self.output_path.unlink()
                build_workbook(
                    self._entities(value), str(self.output_path), mode=CRITIQUE_MODE
                )
                self.assertTrue(self.output_path.exists())

    def test_unknown_mode_is_rejected(self):
        with self.assertRaises(ValueError) as ctx:
            build_workbook(self._entities("Maps cleanly"), str(self.output_path), mode="review")
        self.assertIn("review", str(ctx.exception))


class TestInputFileShape(unittest.TestCase):
    """The mode lives in the persisted JSON, not in a CLI flag.

    The .xlsx is a render and the JSON is the source, so a later session
    regenerating from the JSON alone must get the same validation. A flag
    would let the check pass vacuously the moment someone forgot to pass it.
    """

    def test_load_input_reads_mode_and_entities(self):
        mode, entities = load_input({"mode": "critique", "entities": SAMPLE_ENTITIES})
        self.assertEqual(mode, CRITIQUE_MODE)
        self.assertEqual(entities, SAMPLE_ENTITIES)

    def test_load_input_rejects_bare_list_with_a_migration_hint(self):
        with self.assertRaises(ValueError) as ctx:
            load_input(SAMPLE_ENTITIES)
        message = str(ctx.exception)
        self.assertIn("mode", message)
        self.assertIn("entities", message)

    def test_load_input_rejects_missing_mode(self):
        with self.assertRaises(ValueError) as ctx:
            load_input({"entities": SAMPLE_ENTITIES})
        self.assertIn("mode", str(ctx.exception))

    def test_load_input_rejects_unknown_mode(self):
        with self.assertRaises(ValueError) as ctx:
            load_input({"mode": "as-built", "entities": SAMPLE_ENTITIES})
        self.assertIn("as-built", str(ctx.exception))

    def test_load_input_rejects_missing_entities(self):
        with self.assertRaises(ValueError) as ctx:
            load_input({"mode": "design"})
        self.assertIn("entities", str(ctx.exception))

class TestMergedDocumentationRowHeight(unittest.TestCase):
    """Merged cells are the one place Excel will not autofit, so this row
    keeps an explicit height -- and the chars-per-width estimator that sizes
    it is the same one proven unable to reproduce Excel's metrics.

    Measured against the reference workbook: it uses a flat 30.0pt for every
    merged documentation row, at text lengths from 73 to 337 characters. The
    estimator under-predicted all eight of them (14.5 or 29.0 vs 30.0), which
    means our own 119-123 character docs were rendering at half the reference
    height. The floor removes that whole class of clipping.
    """

    def setUp(self):
        self.output_path = Path(__file__).resolve().parent / "_tmp_docheight.xlsx"
        if self.output_path.exists():
            self.output_path.unlink()

    def tearDown(self):
        if self.output_path.exists():
            self.output_path.unlink()

    def _doc_row_height(self, documentation):
        entities = [
            {
                "name": "E",
                "documentation": documentation,
                "properties": [_property()],
            }
        ]
        build_workbook(entities, str(self.output_path))
        ws = openpyxl.load_workbook(self.output_path)["E"]
        return ws.row_dimensions[2].height

    def test_reference_height_constant_matches_the_measured_reference(self):
        self.assertEqual(REFERENCE_MERGED_DOC_HEIGHT, 30.0)

    def test_short_documentation_gets_the_reference_height_not_one_line(self):
        # 119 chars is the real length of the StudentApplication doc. The bare
        # estimator gave this 14.5pt -- one line -- where the reference gives 30.0.
        height = self._doc_row_height("A" * 119)
        self.assertEqual(height, REFERENCE_MERGED_DOC_HEIGHT)

    def test_two_line_documentation_is_not_under_padded(self):
        height = self._doc_row_height("B" * 337)
        self.assertGreaterEqual(height, REFERENCE_MERGED_DOC_HEIGHT)

    def test_very_long_documentation_grows_past_the_floor(self):
        # The floor is a floor, not a cap -- our docs can exceed anything in
        # the reference workbook.
        height = self._doc_row_height("C" * 3000)
        self.assertGreater(height, REFERENCE_MERGED_DOC_HEIGHT)

    def test_no_documentation_row_is_ever_shorter_than_the_floor(self):
        for length in (1, 73, 109, 127, 243, 262, 337):
            with self.subTest(length=length):
                if self.output_path.exists():
                    self.output_path.unlink()
                self.assertGreaterEqual(
                    self._doc_row_height("D" * length), REFERENCE_MERGED_DOC_HEIGHT
                )


if __name__ == "__main__":
    unittest.main()
